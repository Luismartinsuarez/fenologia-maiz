#!/usr/bin/env python3
"""
Consolida el libro de monitoreo fenológico semanal (una hoja por semana)
en una sola tabla larga: una fila = semana x estación x cultivo.

Uso:
    python consolidar_fenologia.py FENOLOGIA_2021.xlsx
    python consolidar_fenologia.py FENOLOGIA_2021.xlsx -o salida.xlsx --csv
    python consolidar_fenologia.py datos/raw -o datos/processed/FENOLOGIA_consolidado.xlsx --csv
        (varios libros: las semanas repetidas se toman del archivo más reciente)

Requiere: pandas, openpyxl
"""
import argparse
import datetime as dt
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------

# Posición de las columnas en cada hoja (0 = columna A). Se valida contra el
# encabezado de cada hoja antes de leerla.
COLS = {
    "ddrr": 0, "departamento": 1, "provincia": 2, "distrito": 3, "ubigeo": 4,
    "estacion": 5, "cultivo": 6, "variedad": 7, "fecha_siembra": 8,
    "fase": 9, "fecha_inicio_fase": 10, "fecha_observacion": 11, "pct_fase": 12,
    "estado_cultivo": 13, "labores_culturales": 14,
    "fenomeno": 15, "fecha_fenomeno": 16, "pct_fenomeno": 17,
    "plaga": 18, "fecha_plaga": 19, "pct_plaga": 20,
    "observaciones": 21,
}
HEADER_CHECK = {5: "ESTACION", 6: "NOMBRE DE CULTIVO", 9: "FASE FENOLOGICA"}

MESES = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "sep": 9, "oct": 10, "nov": 11, "dic": 12,
}
MESES_NOMBRE = {v: k for k, v in MESES.items() if k != "sep"}

ZONAS = {"COSTA", "SIERRA", "SELVA"}

# Sinónimos de cultivo (se aplican después de quitar tildes y pasar a minúsculas)
CULTIVO_SINONIMOS = {
    "palta": "palto",
    "durazno": "duraznero",
    "habas": "haba",
    "frejol": "frijol",
    "rey grass": "ryegrass",
    "rye grass": "ryegrass",
    "humari": "umari",
    "menbrillero": "membrillero",
    "masgua": "mashua",
}
# Valores de cultivo/fase que en realidad indican que no hay observación
DESCANSO = {"descanso", "terreno en descanso", "cultivo en descanso", "reposo",
            "paralizado", "reposo por inundacion"}
VULNERABLE = {"observador vulnerable"}

# Variantes de nombre de una misma estación -> nombre unificado (se aplica después
# de quitar tildes, pasar a mayúsculas y unificar prefijos). Revisado contra la base:
# las variantes nunca reportan en la misma semana que el nombre unificado, salvo
# MORROPON y SAN MIGUEL, que se solapan unas semanas con parcelas distintas.
ESTACION_TYPOS = {
    # --- errores de digitación y espacios
    "CP-CHUQUBAMBILLA": "CP-CHUQUIBAMBILLA",
    "CP-CAYCAY": "CP-CAYCAY",
    "CO-CAICAY": "CO-CAYCAY",
    "CO-PISAQ": "CO-PISAC",
    "CO-SAN PEDRO HUACARPANA": "CO-SAN PEDRO DE HUACARPANA",
    "CO-SN PEDRO DE HUACARPANA": "CO-SAN PEDRO DE HUACARPANA",
    "CO-CUYOCUYO": "CO-CUYO CUYO",
    "CO-HUARAYAMOHO": "CO-HUARAYA MOHO",
    "CO-ISLALOSUROS": "CO-ISLA LOS UROS",
    "CO-ISLASOTO": "CO-ISLA SOTO",
    "CO-ISLASUANA": "CO-ISLA SUANA",
    "CO-ISLATAQUILE": "CO-ISLA TAQUILE",
    "CO-MAZOCRUZ": "CO-MAZO CRUZ",
    "CO-SANTALUCIA": "CO-SANTA LUCIA",
    "CO-SANTAROSA": "CO-SANTA ROSA",
    "CO-R-DE LA C- ACORA": "CO-R.DE LA C. ACORA",
    "CO-R.DE LA C-ACORA": "CO-R.DE LA C. ACORA",
    "CO-R.DELAC.ACORA": "CO-R.DE LA C. ACORA",
    "CO-RINCON DL CRUZ ACORA": "CO-R.DE LA C. ACORA",
    "HLG-PTE- CALLACAME": "HLG-PTE. CALLACAME",
    "HLG-PTE-CALLACAME": "HLG-PTE. CALLACAME",
    "HLG-PTE- UNOCOLLA": "HLG-PTE. UNOCOLLA",
    "HLG-PTE-UNOCOLLA": "HLG-PTE. UNOCOLLA",
    "PIRAB18:Q20": "PIRA",
    "G": "CO-QUINISTAQUILLAS",           # hoja 2Abr24: misma vid y siembra (10/08/1940)
    # --- cambios de nombre o de prefijo en el tiempo -> nombre más reciente
    "CACHACHI": "PE-CACHACHI",           # hasta dic-2022
    "CHUGUR": "PLU-CHUGUR",              # hasta dic-2022
    "PORCON": "CO-GRANJA PORCON",        # jul-2020
    "GRANJA PORCON": "CO-GRANJA PORCON", # hasta feb-2023
    "SAN MARCOS": "CO-SAN MARCOS",       # hasta feb-2023
    "CO-CUSIPATA": "CUSIPATA",           # hasta set-2023
    "CP-URUBAMBA": "CO-URUBAMBA",        # hasta nov-2021
    "CO-CANCHAN": "CANCHAN",             # Huánuco quitó los prefijos en jun-2024
    "CO-DOS DE MAYO": "DOS DE MAYO",
    "CO-TULUMAYO": "TULUMAYO",
    "CP-HUANUCO": "HUANUCO",
    "CP-PUERTO INCA": "PUERTO INCA",
    "CP-TANANTA": "TANANTA",
    "DIVISORIA": "LA DIVISORIA",
    "PACHACOTO": "NUEVO PACHACOTO",      # 3 semanas justo antes de NUEVO PACHACOTO
    "CO-MORROPON": "MORROPON",           # hasta abr-2025
    "MAP-SAN MIGUEL": "SAN MIGUEL",      # hasta jun-2023
    # --- hoja 4Dic_20: filas sin datos con la DDRR de Puno arrastrada
    "HUAMANI": "CO-HUAMANI",
    "LETRAYOC": "HLG-LETRAYOC",
    "LLAUTA": "HLG-LLAUTA",
    "LOS MOLINOS": "HLG-LOS MOLINOS",
    "SAN BENITO": "CO-SAN BENITO",
}
# Filas de encabezado que no son estaciones (se saltan, como COSTA/SIERRA/SELVA)
ES_ENCABEZADO = re.compile(r"^REGION\s")

# Una DDRR con menos de este % de los registros de una estación se considera un
# error de arrastre y se corrige a la DDRR principal (columna ddrr_norm)
MIN_PCT_DDRR = 10

# ---------------------------------------------------------------------------
# Utilidades de limpieza
# ---------------------------------------------------------------------------

def txt(v):
    """Texto limpio o None."""
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def sin_tildes(s):
    """Quita tildes pero conserva la ñ."""
    s = s.replace("ñ", "\0").replace("Ñ", "\1")
    s = "".join(c for c in unicodedata.normalize("NFKD", s)
                if not unicodedata.combining(c))
    return s.replace("\0", "ñ").replace("\1", "Ñ")


def norm(v):
    """Clave normalizada: minúsculas, sin tildes, espacios simples, sin * finales."""
    s = txt(v)
    if s is None:
        return None
    s = sin_tildes(s).lower().strip(" *.")
    return s or None


def norm_estacion(v):
    s = txt(v)
    if s is None:
        return None
    s = sin_tildes(s).upper()
    # "CO OYOTUN", "CO.ARAPA", "HLM. LAMPA" -> "CO-OYOTUN", "CO-ARAPA", "HLM-LAMPA"
    s = re.sub(r"^(CO|CP|PLU|MAP|HLM|HLG)[\s.\-_]+", r"\1-", s)
    s = re.sub(r"^\*\s*|\s*\(\*\)$", "", s)          # "* CO-TRUJILLO", "CO-CAJABAMBA (*)"
    return ESTACION_TYPOS.get(s, s)


def norm_cultivo(v):
    s = norm(v)
    if s is None or s in {"0", "preparacion"}:
        return None
    return CULTIVO_SINONIMOS.get(s, s)


ANIO_MIN, ANIO_MAX = 1900, dt.date.today().year + 1   # fechas fuera de este rango se consideran errores de digitación


def parse_fecha(v):
    """Devuelve (fecha, texto_original_si_no_se_pudo_convertir).
    Las fechas imposibles (año < ANIO_MIN o > ANIO_MAX) no se convierten y
    quedan como texto para revisión."""
    f, t = _parse_fecha(v)
    if f is not None and not (ANIO_MIN <= f.year <= ANIO_MAX):
        return None, f.strftime("%d/%m/%Y")
    return f, t


def _parse_fecha(v):
    if v is None:
        return None, None
    if isinstance(v, dt.datetime):
        return v.date(), None
    if isinstance(v, dt.date):
        return v, None
    if isinstance(v, (int, float)):
        if 20000 < v < 60000:                      # número de serie de Excel
            return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))), None
        return None, str(int(v)) if float(v).is_integer() else str(v)
    s = txt(v)
    if s is None:
        return None, None
    m = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
    if m:
        d, mth, y = (int(x) for x in m.groups())
        if y < 100:
            y += 2000
        try:
            return dt.date(y, mth, d), None
        except ValueError:
            pass
    return None, s


def parse_pct(cell):
    """Porcentaje en escala 0-100. Corrige valores guardados como fracción con
    formato de porcentaje (0.925 con formato '0.0%' -> 92.5)."""
    v = cell.value if cell is not None else None
    if v is None:
        return None, None
    if isinstance(v, (int, float)):
        fmt = getattr(cell, "number_format", "") or ""
        if "%" in fmt and v <= 1:
            return round(v * 100, 4), None
        return float(v), None
    s = txt(v)
    if s is None:
        return None, None
    nums = re.findall(r"\d+(?:[.,]\d+)?", s)
    if len(nums) == 1 and re.fullmatch(r"[\d.,%\s]+", s):
        return float(nums[0].replace(",", ".")), None
    return None, s                                   # no se pudo; queda el texto


def parse_estado(v):
    """El estado viene como código (1-4) o como texto (Bueno, Regular...)."""
    if v is None:
        return None, None
    if isinstance(v, (int, float)) and float(v).is_integer() and 0 <= v <= 5:
        return int(v), None
    s = norm(v)
    if s is None:
        return None, None
    if re.fullmatch(r"[0-5]", s):
        return int(s), None
    return None, s


def parse_semana_titulo(titulo, semana_n, mes_n, anio):
    """'3ra Semana de Diciembre 15 al 21' -> (inicio, fin)."""
    if not titulo:
        return None, None
    # "15 al 21", "31 Enero al 06", "29 de Julio al 04"
    m = re.search(r"(\d{1,2})\s*(?:de\s+)?(?:[^\W\d_]+\s+)?al\s*(\d{1,2})", titulo, re.I)
    if not m:
        return None, None
    d1, d2 = int(m.group(1)), int(m.group(2))

    def shift(y, mm, k):
        mm += k
        return (y + (mm - 1) // 12, (mm - 1) % 12 + 1)

    if d1 <= d2:                         # semana dentro del mes
        y1, m1 = y2, m2 = anio, mes_n
    elif semana_n == 1:                  # empieza a fines del mes anterior
        (y1, m1), (y2, m2) = shift(anio, mes_n, -1), (anio, mes_n)
    else:                                # termina a inicios del mes siguiente
        (y1, m1), (y2, m2) = (anio, mes_n), shift(anio, mes_n, 1)
    try:
        return dt.date(y1, m1, d1), dt.date(y2, m2, d2)
    except ValueError:
        return None, None


def info_hoja(nombre):
    """'3Dic_20', '3Dic20', '2Novt21' -> (3, 12, 2020)."""
    m = re.fullmatch(r"\s*(\d)\s*([A-Za-z]{3})[A-Za-z]*_?(\d{2})\s*", nombre)
    if not m:
        return None
    sem, mes, yy = m.groups()
    mes_n = MESES.get(mes.lower())
    return (int(sem), mes_n, 2000 + int(yy)) if mes_n else None


def tipo_registro(cultivo_n, fase_n, hay_datos):
    for v in (cultivo_n, fase_n):
        if v in VULNERABLE:
            return "observador vulnerable"
        if v in DESCANSO:
            return "descanso"
    if cultivo_n or hay_datos:
        return "observacion"
    return "sin dato"

# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------

def leer_hoja(ws, avisos, semana_fin_de=None):
    rows = list(ws.iter_rows())
    # Fila de encabezado: la que dice ESTACION (normalmente en la columna F)
    hi = ce = None
    for i, r in enumerate(rows[:15]):
        ce = next((j for j, c in enumerate(r[:8]) if txt(c.value) == "ESTACION"), None)
        if ce is not None:
            hi = i
            break
    if hi is None:
        avisos.append(f"{ws.title}: no se encontró el encabezado, se omite")
        return [], None
    # Algunas hojas no tienen las columnas DEPARTAMENTO..UBIGEO: se desplazan
    # las columnas desde ESTACION y las que faltan quedan vacías.
    desp = ce - COLS["estacion"]
    if desp:
        avisos.append(f"{ws.title}: ESTACION en la columna {ce} (se esperaba "
                      f"{COLS['estacion']}); columnas ajustadas, sin ubicación administrativa")
    cols = {k: (c + desp if c >= COLS["estacion"] else (c if desp == 0 else None))
            for k, c in COLS.items()}
    cols["ddrr"] = COLS["ddrr"]
    for col, esperado in HEADER_CHECK.items():
        val = sin_tildes(txt(rows[hi][col + desp].value) or "").upper()
        if val != esperado:
            avisos.append(f"{ws.title}: col {col + desp} = '{val}', se esperaba '{esperado}'")

    campania = next((txt(r[0].value) for r in rows[:hi]
                     if r and txt(r[0].value) and "MONITOREO" in txt(r[0].value).upper()), None)
    titulo = next((txt(c.value) for r in rows[:hi] for c in r
                   if txt(c.value) and "SEMANA" in txt(c.value).upper()), None)

    semana_fin = semana_fin_de(titulo) if semana_fin_de else None
    ddrr = zona = seccion = None
    regs = []
    for i, r in enumerate(rows[hi + 2:], start=hi + 3):
        cells = list(r) + [None] * (len(COLS) + 2 - len(r))
        cell = lambda k: cells[cols[k]] if cols[k] is not None else None
        val = lambda k: cell(k).value if cell(k) is not None else None

        if txt(val("ddrr")):
            ddrr, zona, seccion = txt(val("ddrr")).upper(), None, None
        est = txt(val("estacion"))
        if not est:
            continue                                # filas vacías o basura sin estación
        if est.upper() in ZONAS:
            zona, seccion = est.upper(), None
            continue
        if est.upper().startswith("CULTIVO DE"):   # separadores (DZ San Martín)
            seccion = est[len("CULTIVO DE"):].strip()
            continue
        if est.upper() in {"NOMBDEP", "ESTACION"} or ES_ENCABEZADO.match(sin_tildes(est).upper()):
            continue

        f_siem, f_siem_txt = parse_fecha(val("fecha_siembra"))
        f_ini, f_ini_txt = parse_fecha(val("fecha_inicio_fase"))
        f_obs, f_obs_txt = parse_fecha(val("fecha_observacion"))
        f_fen, f_fen_txt = parse_fecha(val("fecha_fenomeno"))
        f_pla, f_pla_txt = parse_fecha(val("fecha_plaga"))
        pct, pct_txt = parse_pct(cell("pct_fase"))
        pct_fen, _ = parse_pct(cell("pct_fenomeno"))
        pct_pla, _ = parse_pct(cell("pct_plaga"))
        est_cod, est_txt = parse_estado(val("estado_cultivo"))

        ub = val("ubigeo")
        if isinstance(ub, (int, float)):
            ub = str(int(ub)).zfill(6)
        else:
            ub = txt(ub)
            if ub and ub.isdigit():
                ub = ub.zfill(6)

        cultivo = txt(val("cultivo"))
        cultivo_n = norm_cultivo(cultivo)
        if cultivo_n is None and seccion:
            cultivo_n = norm_cultivo(seccion)
        fase = txt(val("fase"))
        fase_n = norm(fase)
        revisar = []
        if f_ini and semana_fin and f_ini > semana_fin:
            revisar.append("inicio de fase posterior a la semana")
        if f_obs and semana_fin and f_obs > semana_fin + dt.timedelta(days=7):
            revisar.append("observación posterior a la semana")
        if f_siem and f_ini and f_siem > f_ini:
            revisar.append("siembra posterior al inicio de fase")
        if any(t for t in (f_siem_txt, f_ini_txt, f_obs_txt) if t and re.search(r"\d", t) and not re.fullmatch(r"\d{4}", t)):
            revisar.append("fecha no convertible")
        if pct_txt:
            revisar.append("porcentaje no numérico")
        hay_datos = any(x is not None for x in (fase_n, f_ini, pct, est_cod, est_txt))

        regs.append({
            "fila_origen": i,
            "ddrr": ddrr, "zona": zona,
            "departamento": txt(val("departamento")), "provincia": txt(val("provincia")),
            "distrito": txt(val("distrito")), "ubigeo": ub,
            "estacion": est, "estacion_norm": norm_estacion(est),
            "cultivo": cultivo, "cultivo_norm": cultivo_n,
            "variedad": txt(val("variedad")),
            "fecha_siembra": f_siem, "fecha_siembra_texto": f_siem_txt,
            "fase": fase, "fase_norm": fase_n,
            "fecha_inicio_fase": f_ini, "fecha_inicio_fase_texto": f_ini_txt,
            "fecha_observacion": f_obs, "fecha_observacion_texto": f_obs_txt,
            "pct_fase": pct, "pct_fase_texto": pct_txt,
            "estado_codigo": est_cod, "estado_texto": est_txt,
            "labores_culturales": txt(val("labores_culturales")),
            "fenomeno": txt(val("fenomeno")),
            "fecha_fenomeno": f_fen, "fecha_fenomeno_texto": f_fen_txt,
            "pct_fenomeno": pct_fen,
            "plaga": txt(val("plaga")),
            "fecha_plaga": f_pla, "fecha_plaga_texto": f_pla_txt,
            "pct_plaga": pct_pla,
            "observaciones": txt(val("observaciones")),
            "tipo_registro": tipo_registro(cultivo_n, fase_n, hay_datos),
            "revisar": "; ".join(revisar) or None,
        })
    return regs, {"campania": campania, "titulo": titulo}


def consolidar(ruta):
    wb = load_workbook(ruta, read_only=True)
    avisos, todos, resumen = [], [], []
    for ws in wb.worksheets:
        info = info_hoja(ws.title)
        if not info:
            avisos.append(f"{ws.title}: nombre de hoja no reconocido, se omite")
            continue
        sem, mes_n, anio = info
        regs, meta = leer_hoja(
            ws, avisos, lambda t: parse_semana_titulo(t, sem, mes_n, anio)[1])
        if meta is None:
            continue
        ini, fin = parse_semana_titulo(meta["titulo"], sem, mes_n, anio)
        if ini and fin and (fin - ini).days != 6:
            avisos.append(f"{ws.title}: la semana '{meta['titulo']}' no dura 7 días")
        m = re.match(r"\s*(\d)", meta["titulo"] or "")
        if m and int(m.group(1)) != sem:
            avisos.append(f"{ws.title}: el título dice '{meta['titulo']}' "
                          f"(se usa la semana {sem} del nombre de la hoja)")
        cab = {"archivo": Path(ruta).name, "hoja": ws.title.strip(), "anio": anio, "mes": mes_n,
               "mes_nombre": MESES_NOMBRE[mes_n], "semana_mes": sem,
               "semana_inicio": ini, "semana_fin": fin}
        for r in regs:
            todos.append({**cab, **r})
        resumen.append({**cab, "titulo_hoja": meta["titulo"],
                        "campania": meta["campania"], "registros": len(regs),
                        "estaciones": len({r["estacion_norm"] for r in regs}),
                        "observaciones_con_cultivo":
                            sum(r["tipo_registro"] == "observacion" for r in regs)})
    df = pd.DataFrame(todos)
    df = df.sort_values(["anio", "mes", "semana_mes", "ddrr", "zona", "fila_origen"],
                        kind="stable").reset_index(drop=True)
    res = pd.DataFrame(resumen).sort_values(["anio", "mes", "semana_mes"])
    return df, res, avisos


def asignar_ddrr(df):
    """ddrr_norm: DDRR de la estación, corrigiendo filas con la DDRR mal arrastrada.
    Las DDRR con menos de MIN_PCT_DDRR % de los registros de la estación se
    reemplazan por la principal. Si quedan dos o más (homónimos: lugares distintos
    con el mismo nombre, p. ej. CO-SAN PABLO en Cajamarca y San Martín) se respetan."""
    n = df.groupby(["estacion_norm", "ddrr"]).size().rename("n").reset_index()
    n["pct"] = 100 * n["n"] / n.groupby("estacion_norm")["n"].transform("sum")
    principal = n.sort_values("n").groupby("estacion_norm")["ddrr"].last()
    validas = set(zip(n.loc[n["pct"] >= MIN_PCT_DDRR, "estacion_norm"],
                      n.loc[n["pct"] >= MIN_PCT_DDRR, "ddrr"]))
    ok = [k in validas for k in zip(df["estacion_norm"], df["ddrr"])]
    df["ddrr_norm"] = df["ddrr"].where(ok, df["estacion_norm"].map(principal))
    cambiadas = (df["ddrr_norm"] != df["ddrr"]) & df["ddrr"].notna()
    homon = n[n["pct"] >= MIN_PCT_DDRR].groupby("estacion_norm").size().loc[lambda s: s > 1]
    avisos = [f"ddrr_norm: {int(cambiadas.sum())} filas con la DDRR corregida a la principal de su estación"]
    avisos += [f"homónimo: {e} aparece en varias DDRR "
               f"({', '.join(sorted(n.loc[(n['estacion_norm'] == e) & (n['pct'] >= MIN_PCT_DDRR), 'ddrr']))})"
               for e in homon.index]
    pos = df.columns.get_loc("ddrr") + 1
    df.insert(pos, "ddrr_norm", df.pop("ddrr_norm"))
    return df, avisos


CLAVE_SEMANA = ["anio", "mes", "semana_mes"]
COLS_COMPARAR = ["estacion_norm", "cultivo_norm", "variedad", "fecha_siembra",
                 "fase_norm", "fecha_inicio_fase", "pct_fase", "estado_codigo",
                 "estado_texto", "fenomeno", "plaga"]


def _firma(g):
    """Contenido de una semana como multiconjunto de filas comparables."""
    return sorted(tuple("" if pd.isna(v) else str(v) for v in fila)
                  for fila in g[COLS_COMPARAR].itertuples(index=False))


def consolidar_varios(rutas):
    """Consolida varios libros. Una misma semana puede venir en más de un
    archivo; se conserva la versión del archivo más reciente (el que empieza
    más tarde). En la práctica, la última hoja de un libro viejo suele ser una
    copia sin actualizar de la semana anterior, y el libro de la campaña
    siguiente trae los datos reales."""
    dfs, ress, avisos = [], [], []
    for ruta in rutas:
        df, res, av = consolidar(ruta)
        nombre = Path(ruta).name
        avisos += [f"{nombre} | {a}" for a in av]
        print(f"  {nombre}: {len(df):,} registros, {len(res)} hojas", file=sys.stderr)
        dfs.append(df)
        ress.append(res)
    res = pd.concat(ress, ignore_index=True)

    # Prioridad: primera semana más tardía > más hojas > nombre más corto/alfabético
    # (FENOLOGIA_2324 antes que la copia FENOLOGIA_23241111)
    res["_orden"] = res["anio"] * 1000 + res["mes"] * 10 + res["semana_mes"]
    pri = (res.groupby("archivo")
              .agg(primera=("_orden", "min"), hojas=("hoja", "size"))
              .reset_index()
              .sort_values(["primera", "hojas", "archivo"], ascending=[False, False, True]))
    rango = {a: i for i, a in enumerate(pri["archivo"])}       # 0 = preferido
    res["_rango"] = res["archivo"].map(rango)
    res = res.sort_values(CLAVE_SEMANA + ["_rango"])
    elegido = res.groupby(CLAVE_SEMANA)["archivo"].transform("first")

    df = pd.concat(dfs, ignore_index=True)
    grupos = {k: g for k, g in df.groupby(CLAVE_SEMANA + ["archivo"])}
    dup = []
    for k, g in res.groupby(CLAVE_SEMANA):
        if len(g) < 2:
            continue
        usado = g["archivo"].iloc[0]
        f_usado = _firma(grupos.get((*k, usado), df.iloc[:0]))
        for a in g["archivo"].iloc[1:]:
            f_otro = _firma(grupos.get((*k, a), df.iloc[:0]))
            dup.append({"anio": k[0], "mes": k[1], "semana_mes": k[2],
                        "hoja": g["hoja"].iloc[0], "archivo_usado": usado,
                        "archivo_descartado": a,
                        "registros_usado": len(f_usado), "registros_descartado": len(f_otro),
                        "contenido_identico": f_usado == f_otro,
                        "filas_distintas": sum((Counter(f_usado) - Counter(f_otro)).values())})
    dup = pd.DataFrame(dup)
    if len(dup):
        dif = dup[~dup["contenido_identico"]]
        for _, d in dif.iterrows():
            avisos.append(f"{d['hoja']} ({d['anio']}): {d['filas_distintas']} filas difieren entre "
                          f"{d['archivo_usado']} (usado) y {d['archivo_descartado']} (descartado)")

    res["usado"] = res["archivo"] == elegido
    keep = res.loc[res["usado"], CLAVE_SEMANA + ["archivo"]]
    df = df.merge(keep, on=CLAVE_SEMANA + ["archivo"], how="inner")
    df = df.sort_values(CLAVE_SEMANA + ["ddrr", "zona", "fila_origen"],
                        kind="stable").reset_index(drop=True)
    res = res.drop(columns=["_rango", "_orden"]).reset_index(drop=True)
    return df, res, dup, avisos

# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------

DESCRIPCION = [
    ("archivo", "Libro Excel de origen"),
    ("hoja", "Nombre de la hoja de origen (p. ej. 3Dic_20 = 3ra semana de diciembre 2020)"),
    ("anio / mes / semana_mes", "Tomados del nombre de la hoja"),
    ("semana_inicio / semana_fin", "Fechas de la semana, calculadas del título de la hoja ('15 al 21')"),
    ("fila_origen", "Fila en la hoja original, para rastrear cualquier dato"),
    ("ddrr / zona", "Dirección zonal y región natural (rellenadas hacia abajo desde las filas de bloque)"),
    ("ubigeo", "Código de 6 dígitos como texto (conserva el cero inicial)"),
    ("estacion_norm", "Nombre de estación sin tildes, con prefijos unificados (CO-, CP-, PLU-, MAP-, HLM-, HLG-) y variantes del mismo nombre unificadas (ESTACION_TYPOS en el script; en cambios de nombre se usa el más reciente). 'estacion' conserva el texto original"),
    ("ddrr_norm", "DDRR de la estación. Corrige filas con la DDRR mal arrastrada (menos del 10 % de los registros de la estación). Para agrupar por estación, usar ddrr_norm + estacion_norm"),
    ("cultivo_norm", "Cultivo en minúsculas, sin tildes y con sinónimos unificados (palta→palto, habas→haba...)"),
    ("fase_norm", "Fase en minúsculas y sin tildes (no se agrupan sinónimos de fase)"),
    ("fecha_*", "Fecha convertida. Si no se pudo convertir, el valor original queda en la columna *_texto"),
    ("pct_*", "Porcentaje en escala 0-100 (valores guardados como fracción con formato % se multiplicaron por 100)"),
    ("estado_codigo / estado_texto", "El original mezcla códigos 1-4 con texto; se separan. El significado de los códigos no está en el archivo"),
    ("observaciones", "Columna V de la hoja original (notas: 'NO REPORTA', 'Situación de vulnerabilidad', etc.)"),
    ("tipo_registro", "observacion | descanso | observador vulnerable | sin dato"),
    ("revisar", "Inconsistencias detectadas (fechas imposibles o fuera de orden, % no numérico). Vacío = sin alertas"),
    ("semanas.usado", "Si una semana viene en varios archivos se usa la del archivo más reciente (el que empieza más tarde); las demás quedan con usado = False"),
    ("duplicados", "Cada semana repetida: archivo usado, archivo descartado y si el contenido era idéntico"),
]


def escribir_excel(df, res, avisos, salida, dup=None):
    with pd.ExcelWriter(salida, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="consolidado", index=False)
        res.to_excel(xw, sheet_name="semanas", index=False)
        if dup is not None:
            (dup if len(dup) else pd.DataFrame({"aviso": ["(sin semanas duplicadas)"]})
             ).to_excel(xw, sheet_name="duplicados", index=False)
        pd.DataFrame(DESCRIPCION, columns=["columna", "descripcion"]).to_excel(
            xw, sheet_name="leeme", index=False)
        pd.DataFrame({"aviso": avisos or ["(sin avisos)"]}).to_excel(
            xw, sheet_name="avisos", index=False)

        head_fill = PatternFill("solid", fgColor="1F4E78")
        for ws in xw.book.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    c.font = Font(name="Arial", size=10)
            for c in ws[1]:
                c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                c.fill = head_fill
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), 1):
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 9), 45)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, (dt.date, dt.datetime)):
                        c.number_format = "DD/MM/YYYY"
        xw.book["leeme"].column_dimensions["B"].width = 100


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("entrada", nargs="+",
                    help="Uno o más libros Excel (o carpetas con libros) con una hoja por semana")
    ap.add_argument("-o", "--salida", help="Archivo Excel de salida")
    ap.add_argument("--csv", action="store_true", help="Guardar también un CSV (UTF-8)")
    a = ap.parse_args()

    rutas = []
    for e in map(Path, a.entrada):
        rutas += sorted(p for p in e.glob("*.xlsx") if not p.name.startswith("~$")
                        and not p.name.startswith(".~")) if e.is_dir() else [e]
    if not rutas:
        ap.error("no se encontraron archivos .xlsx")

    if len(rutas) == 1:
        df, res, avisos = consolidar(rutas[0])
        dup = None
        defecto = rutas[0].with_name(rutas[0].stem + "_consolidado.xlsx")
    else:
        df, res, dup, avisos = consolidar_varios(rutas)
        defecto = Path("FENOLOGIA_consolidado.xlsx")
    df, av = asignar_ddrr(df)
    avisos += av
    salida = Path(a.salida) if a.salida else defecto
    escribir_excel(df, res, avisos, salida, dup)
    if a.csv:
        df.to_csv(salida.with_suffix(".csv"), index=False, encoding="utf-8-sig")

    print(f"{len(df):,} registros de {res['usado'].sum() if dup is not None else len(res)} "
          f"semanas ({len(rutas)} archivos) -> {salida}")
    print(df["tipo_registro"].value_counts().to_string())
    for x in avisos:
        print("AVISO:", x, file=sys.stderr)


if __name__ == "__main__":
    main()
