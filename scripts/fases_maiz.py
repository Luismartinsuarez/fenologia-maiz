#!/usr/bin/env python3
"""
Fechas de las fases fenológicas del maíz por estación y variedad.

Parte de la base consolidada (salida de consolidar_fenologia.py) y arma una
fila por campaña: estación + fecha de siembra, con la fecha de inicio de cada
fase y los días desde la siembra.

Uso:
    python fases_maiz.py datos/processed/FENOLOGIA_consolidado.csv
    python fases_maiz.py datos/processed/FENOLOGIA_consolidado.csv -o fases_maiz.xlsx

Requiere: pandas, openpyxl
"""
import argparse
import re
import sys
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from consolidar_fenologia import norm  # noqa: E402

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------

# Fases canónicas en orden de desarrollo
FASES = ["siembra", "emergencia", "aparicion de hojas", "panoja", "espiga",
         "maduracion lechosa", "maduracion pastosa", "maduracion cornea", "cosecha"]
ORDEN = {f: i for i, f in enumerate(FASES)}

# Se prueban de la más avanzada a la menos avanzada: si un texto menciona
# varias fases ("aparicion de 14 hojas panoja") se queda la más avanzada.
REGLAS_FASE = [
    ("cosecha", r"cosech|^corte"),
    ("maduracion cornea", r"corn|conea"),
    ("maduracion pastosa", r"pasto|patosa"),
    ("maduracion lechosa", r"lech|lehosa"),
    ("espiga", r"espiga|esgiga"),
    ("panoja", r"panoja"),
    ("aparicion de hojas", r"hoja"),
    ("emergencia", r"emer|plantula"),
    ("siembra", r"siembra"),
]
# Fases que no se pueden ubicar en el esquema (se listan, no van a las fechas)
SIN_ESQUEMA = {"maduracion": "maduracion (sin especificar)",
               "floracion": "floracion (sin especificar)"}

NUM_PALABRA = [  # (prefijo, número); los más largos primero
    ("dieciseis", 16), ("veitiun", 21), ("veintiun", 21), ("catorce", 14),
    ("quince", 15), ("trece", 13), ("doce", 12), ("once", 11), ("undecim", 11),
    ("diez", 10), ("decim", 10), ("nueve", 9), ("nuene", 9), ("noven", 9),
    ("ocho", 8), ("octav", 8), ("siete", 7), ("sietes", 7), ("septim", 7),
    ("setim", 7), ("stiem", 7), ("seis", 6), ("sext", 6), ("cinco", 5),
    ("quint", 5), ("cuatro", 4), ("cuart", 4), ("tres", 3), ("tercer", 3),
    ("dos", 2), ("segund", 2), ("primer", 1),
]

VARIEDAD_SINONIMOS = {
    "san jeronimo punta roja": "san geronimo punta roja",
    "marginal 28t": "marginal 28", "marginal 28 t": "marginal 28",
    "dekal": "dekalb", "blanca": "blanco", "amilaceo blanco": "blanco amilaceo",
    "mibrido bm 3063": "hibrido bm 3063",
    "compusto huancavelicano": "compuesto huancavelicano",
    "amarillo comun": "amarillo", "oro amarillo": "amarillo oro",
    # Errores de escritura y formas de una misma variedad (revisadas por estación)
    "atlas-777": "atlas 777", "mujo de atlas 777": "atlas 777",
    "marginla 28": "marginal 28", "marginal": "marginal 28",
    "dekal b": "dekalb", "dkal 399": "dekalb 399", "amarillo de kalb 7500": "dekalb 7500",
    "dekal 7508": "dekalb 7508", "dekal-7508": "dekalb 7508",
    "dekal 7088": "dekalb 7088", "dekal-7088": "dekalb 7088",
    "advance": "advanta", "adv 9139": "advanta 9139",
    "adv 9789": "advanta 9789", "amarillo adv 9789": "advanta 9789",
    "amarillo hibridonk252": "nk 252",
    "hr960": "inia 619 megahibrido", "megahibrido he960": "inia 619 megahibrido",
    "megahibrido hr960": "inia 619 megahibrido", "inia 619": "inia 619 megahibrido",
    "hibrido duro": "amarillo duro hibrido", "amarillo hibrido duro": "amarillo duro hibrido",
    "akira inia 628": "akira", "amarillo duro inia 626-akira": "akira",
    "chuska": "inia 617 chuska", "inia 617-chusca": "inia 617 chuska",
    "nutrimax": "nutrimaiz",
    "san jeronimo": "san geronimo punta roja", "san geronimo inia": "san geronimo punta roja",
    "punta roja": "san geronimo punta roja",
    # Blanco Urubamba = Blanco Gigante Cusco
    "blanco gigante": "blanco urubamba", "gigante cusco": "blanco urubamba",
    "blanco gigante de cusco": "blanco urubamba", "blanco gigante cusco": "blanco urubamba",
    "cusco urubamba": "blanco urubamba",
    "morado de canta": "morado canteño", "canta": "morado canteño", "morado de cornea": "morado canteño",
    "cusqueado inia": "cusqueado", "8 rayas (cusqueado)": "cusqueado",
    "blanco cusqueado": "cusqueado", "blanco 8 rayas": "cusqueado",
    "opaco": "opaco mal paso", "astillo": "astilla", "cabanita blanco": "cabanita",
    "imperial blanco": "blanco imperial", "duro amarillo perejil": "perejil",
    "amilaceo blanco de la zona": "blanco amilaceo de la zona", "correinte": "corriente",
    "comun": "comun de la zona", "amilaceo chaucha": "chaucha", "comun imperial": "imperial",
    "blanco criollo": "criollo",
    # No son variedades: parcela en descanso
    "descanso": None, "descando": None,
}

# Sinónimos que solo valen en una estación: el nombre corto o la marca sin número
# se asigna a la variedad que se reporta en esa misma estación
VARIEDAD_POR_ESTACION = {
    ("CO-CALLANCAS", "amarillo de la zona"): "amarillo",
    ("COLCABAMBA", "astilla"): "astilla huancavelicano",
    ("HUASACACHE", "morado"): "morado canteño",
    ("PACARAN", "morado"): "morado canteño",
    ("PLU-LIVES", "amarillo duro hibrido"): "duro hibrido marginal",
    ("CO-OYOTUN", "dekalb"): "dekalb 7088",
    ("CP-TALLA", "dekalb"): "dekalb (sin especificar)",   # en Talla hay 7088 y 7508
    ("CO-ALAO", "advanta"): "advanta 9139",
    ("CAMAY", "advanta"): "advanta (sin especificar)",
    ("CO-ALAO", "hibrido"): "hibrido (sin especificar)",
}

# Nombres que agrupan maíces locales distintos: sirven para analizar cada estación
# por separado, no para requerimientos de una variedad entre estaciones
VARIEDADES_GENERICAS = {
    "blanco", "amarillo", "de la zona", "blanco amilaceo", "blanco amilaceo de la zona",
    "blanco de la zona", "amilaceo", "nativo", "criollo", "blanco criollo de la zona",
    "comun de la zona", "corriente", "amarillo de la zona",
}

GAP_DIAS = 70     # sin reportes por más de esto -> otra campaña
MAX_DIAS = 365    # fase a más de esto de la siembra -> se marca para revisar

# ---------------------------------------------------------------------------
# Normalización
# ---------------------------------------------------------------------------

def fase_maiz(fase_norm):
    """Texto de fase -> (fase canónica o None, número de hojas o None)."""
    if not isinstance(fase_norm, str):
        return None, None
    s = re.sub(r"\s+", " ", fase_norm)
    for fase, patron in REGLAS_FASE:
        if re.search(patron, s):
            return fase, (num_hojas(s) if fase == "aparicion de hojas" else None)
    for clave, fase in SIN_ESQUEMA.items():
        if clave in s:
            return fase, None
    return None, None


def num_hojas(s):
    """Mayor número de hojas mencionado: '8-10 hojas' -> 10, 'novena hoja' -> 9."""
    nums = [int(n) for n in re.findall(r"\d+", s) if 0 < int(n) <= 30]
    if nums:
        return max(nums)
    for tok in re.findall(r"[a-zñ]+", s):
        for pref, n in NUM_PALABRA:
            if tok.startswith(pref):
                nums.append(n)
                break
    return max(nums) if nums else None


def norm_variedad(v, estacion=None):
    s = None if pd.isna(v) else norm(v)
    if s is None:
        return None
    s = re.sub(r"^maiz\s+", "", s)
    s = re.sub(r"\s*-\s*", "-", s)
    s = VARIEDAD_SINONIMOS.get(s, s)
    return VARIEDAD_POR_ESTACION.get((estacion, s), s)


def moda(serie):
    """Valor más frecuente; en empate, el menor (la fecha más temprana)."""
    v = serie.dropna()
    if v.empty:
        return None
    c = Counter(v)
    top = max(c.values())
    return min(k for k, n in c.items() if n == top)

# ---------------------------------------------------------------------------
# Cálculo
# ---------------------------------------------------------------------------

def leer_base(ruta):
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".csv":
        df = pd.read_csv(ruta, dtype={"ubigeo": str}, low_memory=False)
    else:
        df = pd.read_excel(ruta, sheet_name="consolidado", dtype={"ubigeo": str})
    for c in ["semana_inicio", "fecha_siembra", "fecha_inicio_fase"]:
        df[c] = pd.to_datetime(df[c], errors="coerce")
    if "ddrr_norm" in df.columns:         # DDRR corregida (bases generadas desde oct-2026)
        df["ddrr"] = df["ddrr_norm"]
    return df


def ubicacion_estaciones(df):
    """Datos administrativos más frecuentes de cada estación + DDRR (algunas hojas no
    los traen; la DDRR separa homónimos como CO-SAN PABLO de Cajamarca y San Martín)."""
    cols = ["zona", "departamento", "provincia", "distrito", "ubigeo"]
    return df.groupby(["estacion_norm", "ddrr"])[cols].agg(moda).reset_index()


def asignar_ciclos(m):
    """Campaña = estación + fecha de siembra, partida si hay un hueco de reportes
    mayor a GAP_DIAS. Los registros sin fecha de siembra se asignan a la campaña
    de la misma estación cuyo período los contiene."""
    m = m.sort_values(["estacion_norm", "fecha_siembra", "semana_inicio"]).copy()
    con = m[m["fecha_siembra"].notna()].copy()
    salto = con.groupby(["estacion_norm", "fecha_siembra"])["semana_inicio"].diff().dt.days
    con["_tramo"] = (salto > GAP_DIAS).groupby(
        [con["estacion_norm"], con["fecha_siembra"]]).cumsum()
    claves = con[["estacion_norm", "fecha_siembra", "_tramo"]].drop_duplicates()
    claves["ciclo_id"] = range(1, len(claves) + 1)
    con = con.merge(claves, on=["estacion_norm", "fecha_siembra", "_tramo"])

    periodo = con.groupby(["ciclo_id", "estacion_norm"])["semana_inicio"].agg(["min", "max"]).reset_index()
    sin = m[m["fecha_siembra"].isna()].copy()
    asig = []
    for idx, r in sin.iterrows():
        p = periodo[(periodo["estacion_norm"] == r["estacion_norm"])
                    & (periodo["min"] <= r["semana_inicio"]) & (periodo["max"] >= r["semana_inicio"])]
        asig.append(p["ciclo_id"].iloc[0] if len(p) == 1 else None)
    sin["ciclo_id"] = asig
    sin = sin[sin["ciclo_id"].notna()]
    return pd.concat([con.drop(columns="_tramo"), sin], ignore_index=True), len(asig) - len(sin)


def fechas_fase(g):
    """Fecha de inicio de una fase en una campaña, a partir de los reportes semanales."""
    f = moda(g["fecha_inicio_fase"])
    fuente = "fecha de inicio reportada"
    if f is None:
        f, fuente = g["semana_inicio"].min(), "primera semana en que se reporta"
    distintas = g["fecha_inicio_fase"].dropna().nunique()
    return pd.Series({
        "fecha_inicio": f, "fuente": fuente,
        "fechas_reportadas_distintas": distintas,
        "fecha_min_reportada": g["fecha_inicio_fase"].min(),
        "fecha_max_reportada": g["fecha_inicio_fase"].max(),
        "primera_semana": g["semana_inicio"].min(),
        "ultima_semana": g["semana_inicio"].max(),
        "semanas_reportada": g["semana_inicio"].nunique(),
        "pct_max": g["pct_fase"].max(),
    })


def calcular(df):
    es_maiz = df["cultivo_norm"].fillna("").str.contains(r"\bmaiz\b|choclo")
    m = df[es_maiz & (df["tipo_registro"] == "observacion")].copy()
    m[["fase_maiz", "hojas"]] = m["fase_norm"].apply(lambda s: pd.Series(fase_maiz(s)))
    m["variedad_norm"] = [norm_variedad(v, e) for v, e in zip(m["variedad"], m["estacion_norm"])]

    catalogo = (m.groupby(["fase_norm", "fase_maiz", "hojas"], dropna=False).size()
                  .rename("registros").reset_index()
                  .sort_values(["fase_maiz", "hojas", "registros"], ascending=[True, True, False]))

    m, sin_ciclo = asignar_ciclos(m)
    m["ciclo_id"] = m["ciclo_id"].astype(int)

    # --- datos de cada campaña
    cic = m.groupby("ciclo_id").agg(
        estacion=("estacion_norm", "first"),
        ddrr=("ddrr", moda),
        cultivo=("cultivo_norm", moda),
        variedad=("variedad_norm", moda),
        variedades_reportadas=("variedad_norm", lambda s: " | ".join(sorted(s.dropna().unique()))),
        fecha_siembra=("fecha_siembra", moda),
        primera_semana=("semana_inicio", "min"),
        ultima_semana=("semana_inicio", "max"),
        semanas_con_reporte=("semana_inicio", "nunique"),
    ).reset_index()
    cic = cic.merge(ubicacion_estaciones(df), left_on=["estacion", "ddrr"],
                    right_on=["estacion_norm", "ddrr"], how="left").drop(columns="estacion_norm")

    # --- fecha de cada fase (largo); las hojas se detallan por número
    v = m[m["fase_maiz"].isin(FASES)]
    largo = v.groupby(["ciclo_id", "fase_maiz"]).apply(fechas_fase, include_groups=False).reset_index()
    hojas = (v[v["fase_maiz"] == "aparicion de hojas"]
             .groupby(["ciclo_id"])["hojas"].max().rename("hojas_max").reset_index())
    largo = largo.merge(cic[["ciclo_id", "fecha_siembra"]], on="ciclo_id")
    largo["dias_desde_siembra"] = (largo["fecha_inicio"] - largo["fecha_siembra"]).dt.days
    largo["orden_fase"] = largo["fase_maiz"].map(ORDEN)
    largo = largo.rename(columns={"fase_maiz": "fase"}).sort_values(["ciclo_id", "orden_fase"])

    # Revisión: fase antes de la siembra, demasiado lejos, o fuera de orden
    largo["revisar"] = None
    largo.loc[largo["dias_desde_siembra"] < 0, "revisar"] = "fase anterior a la siembra"
    largo.loc[largo["dias_desde_siembra"] > MAX_DIAS, "revisar"] = f"más de {MAX_DIAS} días desde la siembra"
    previa = largo.groupby("ciclo_id")["fecha_inicio"].cummax().groupby(largo["ciclo_id"]).shift()
    fuera = largo["fecha_inicio"] < previa
    largo.loc[fuera & largo["revisar"].isna(), "revisar"] = "fecha anterior a una fase previa"

    # --- una fila por campaña (ancho)
    fe = largo.pivot(index="ciclo_id", columns="fase", values="fecha_inicio")
    di = largo.pivot(index="ciclo_id", columns="fase", values="dias_desde_siembra")
    # La siembra ya está en fecha_siembra (columna del registro), no se repite
    fases_pres = [f for f in FASES if f in fe.columns and f != "siembra"]
    fe = fe[fases_pres].add_prefix("fecha_")
    di = di[fases_pres].add_prefix("dias_")
    rev = largo.dropna(subset=["revisar"]).groupby("ciclo_id").apply(
        lambda g: "; ".join(f"{f}: {r}" for f, r in zip(g["fase"], g["revisar"])),
        include_groups=False).rename("revisar")
    ancho = (cic.merge(hojas, on="ciclo_id", how="left")
                .merge(fe, on="ciclo_id", how="left").merge(di, on="ciclo_id", how="left")
                .merge(rev, on="ciclo_id", how="left"))
    ancho["variedad_generica"] = ancho["variedad"].isin(VARIEDADES_GENERICAS)
    ancho["fases_con_fecha"] = largo.groupby("ciclo_id").size().reindex(ancho["ciclo_id"]).values
    ancho = ancho.sort_values(["ddrr", "estacion", "fecha_siembra"]).reset_index(drop=True)
    cols_id = ["ciclo_id", "ddrr", "zona", "departamento", "provincia", "distrito", "ubigeo",
               "estacion", "cultivo", "variedad", "variedad_generica", "variedades_reportadas", "fecha_siembra",
               "primera_semana", "ultima_semana", "semanas_con_reporte", "fases_con_fecha", "hojas_max"]
    ancho = ancho[cols_id + [c for c in ancho.columns if c not in cols_id and c != "revisar"] + ["revisar"]]

    largo = largo.merge(cic[["ciclo_id", "estacion", "variedad"]], on="ciclo_id")
    largo = largo[["ciclo_id", "estacion", "variedad", "fecha_siembra", "orden_fase", "fase",
                   "fecha_inicio", "dias_desde_siembra", "fuente", "fechas_reportadas_distintas",
                   "fecha_min_reportada", "fecha_max_reportada", "primera_semana", "ultima_semana",
                   "semanas_reportada", "pct_max", "revisar"]]

    # --- resúmenes: mediana de días desde la siembra, solo fases sin alertas
    ok = largo[largo["revisar"].isna() & (largo["fase"] != "siembra")]

    def resumen(por):
        r = ok.merge(ancho[["ciclo_id"] + [p for p in por if p not in ok.columns]], on="ciclo_id")
        t = r.pivot_table(index=por, columns="fase", values="dias_desde_siembra",
                          aggfunc="median")
        t = t[[f for f in FASES if f in t.columns]].add_prefix("mediana_dias_").round(0)
        n = ancho.groupby(por)["ciclo_id"].nunique().rename("campanias")
        return n.to_frame().join(t).reset_index().sort_values(por)

    res_var = resumen(["variedad"])
    res_est = resumen(["ddrr", "estacion", "variedad"])
    return ancho, largo, res_var, res_est, catalogo, sin_ciclo

# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------

LEEME = [
    ("campanias", "Una fila por campaña: estación + fecha de siembra. Si una estación repite la misma fecha de siembra pero hay más de %d días sin reportes, se cuenta como otra campaña" % GAP_DIAS),
    ("variedad", "Variedad más reportada en la campaña (minúsculas, sin tildes, sinónimos unificados). variedades_reportadas lista todas"),
    ("variedad_generica", "True si el nombre agrupa maíces locales distintos (blanco, amarillo, de la zona…): usar solo dentro de cada estación"),
    ("fecha_<fase>", "Fecha de inicio de la fase: la fecha de inicio más repetida en los reportes semanales. Si nunca se reportó la fecha, se usa la primera semana en que aparece la fase"),
    ("dias_<fase>", "Días desde la siembra hasta el inicio de la fase"),
    ("hojas_max", "Mayor número de hojas reportado en la fase 'aparicion de hojas'"),
    ("fases", "Una fila por campaña y fase, con detalle de cómo se obtuvo la fecha (fechas distintas reportadas, primera y última semana, % máximo)"),
    ("revisar", "Fase anterior a la siembra, a más de %d días de la siembra, o con fecha anterior a una fase previa" % MAX_DIAS),
    ("resumen_*", "Mediana de días desde la siembra por variedad y por estación+variedad. Excluye las fases marcadas en 'revisar'"),
    ("catalogo_fases", "Cómo se clasificó cada texto de fase original. Revisar aquí si alguna fase quedó mal asignada"),
    ("fases canónicas", " → ".join(FASES)),
]


def escribir(salida, hojas):
    with pd.ExcelWriter(salida, engine="openpyxl") as xw:
        for nombre, d in hojas.items():
            d.to_excel(xw, sheet_name=nombre, index=False)
        head = PatternFill("solid", fgColor="1F4E78")
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), 1):
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 9), 40)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if hasattr(c.value, "year"):
                        c.number_format = "DD/MM/YYYY"
        xw.book["leeme"].column_dimensions["B"].width = 110


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("entrada", help="Base consolidada (.csv o .xlsx de consolidar_fenologia.py)")
    ap.add_argument("-o", "--salida", help="Excel de salida (por defecto fases_maiz.xlsx junto a la entrada)")
    a = ap.parse_args()

    entrada = Path(a.entrada)
    salida = Path(a.salida) if a.salida else entrada.with_name("fases_maiz.xlsx")
    ancho, largo, res_var, res_est, catalogo, sin_ciclo = calcular(leer_base(entrada))
    escribir(salida, {
        "campanias": ancho, "fases": largo, "resumen_variedad": res_var,
        "resumen_estacion": res_est, "catalogo_fases": catalogo,
        "leeme": pd.DataFrame(LEEME, columns=["hoja / columna", "descripcion"]),
    })
    print(f"{len(ancho)} campañas de maíz, {ancho['estacion'].nunique()} estaciones, "
          f"{ancho['variedad'].nunique()} variedades -> {salida}")
    print(f"fases con fecha: {len(largo)}  | marcadas para revisar: {largo['revisar'].notna().sum()}")
    if sin_ciclo:
        print(f"{sin_ciclo} registros sin fecha de siembra no se pudieron asignar a una campaña")


if __name__ == "__main__":
    main()
